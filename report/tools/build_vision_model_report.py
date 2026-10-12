"""비전 AI 모델 5종(YOLOE · SAM 2 · SigLIP 2 · EasyOCR · MediaPipe Hands) 분석 리포트(docx)를 만든다.

수치는 모두 각 모델의 공식 문서·논문·저장소에 공개된 제원이다 (출처는 리포트 끝에 정리).
직접 측정한 값은 쓰지 않는다.

사용법:
    py report/tools/build_vision_model_report.py
"""
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TOOLS = Path(__file__).resolve().parent
REPORT = TOOLS.parent
FIG = TOOLS / "figures"
OUT = REPORT / "비전_AI_모델_5종_분석_리포트.docx"

FONT = "맑은 고딕"
INK, INK2, GRID = "#0b0b0b", "#52514e", "#e4e3df"
SERIES = "#2a78d6"   # 단일 계열: 기본 팔레트 1번(파랑)

# ---------------------------------------------------------------------------
# 공식 제원 (출처: 각 모델 공식 README·문서·논문)
# ---------------------------------------------------------------------------
# YOLOE-11, 텍스트 프롬프트 기준 — THU-MIG/yoloe README (LVIS minival Fixed AP, T4 TensorRT FPS)
YOLOE11 = [("S", 10, 27.5, 301.2), ("M", 21, 33.0, 168.3), ("L", 26, 35.2, 130.5)]   # (크기, 파라미터 M, AP, FPS)
# SAM 2.1 — facebookresearch/sam2 README (A100 FPS, SA-V test J&F)
SAM21 = [("tiny", 38.9, 76.5, 91.2), ("small", 46.0, 76.6, 84.8), ("base+", 80.8, 78.2, 64.1), ("large", 224.4, 79.5, 39.5)]
# SigLIP 2 — 논문 Table 1, 256px (비전 인코더 파라미터 M, ImageNet-1k 제로샷 %)
SIGLIP2 = [("B/16", 86, 79.1), ("L/16", 303, 82.5), ("So400m/16", 400, 83.4), ("g/16", 1000, 84.5)]

# ---------------------------------------------------------------------------
# 그림: 모델별 '크기를 키우면 무엇을 얻고 무엇을 잃나' (작은 다중 그래프, 패널마다 축이 다르다)
# ---------------------------------------------------------------------------
plt.rcParams.update({
    "font.family": ["Malgun Gothic", "DejaVu Sans"], "axes.unicode_minus": False,
    "font.size": 9.5, "axes.edgecolor": GRID, "axes.labelcolor": INK2,
    "xtick.color": INK2, "ytick.color": INK2, "axes.spines.top": False, "axes.spines.right": False,
})


def panel(ax, xs, ys, labels, title, xlabel, ylabel, logx=False, invert_x=False, offsets=None):
    """offsets = {점 번호: (x, y, 정렬)} 라벨이 겹치는 점만 위치를 따로 지정"""
    ax.plot(xs, ys, color=SERIES, linewidth=2, zorder=2)
    ax.scatter(xs, ys, s=64, color=SERIES, edgecolor="white", linewidth=2, zorder=3)
    for i, (x, y, t) in enumerate(zip(xs, ys, labels)):
        dx, dy, ha = (offsets or {}).get(i, (0, 9, "center"))
        ax.annotate(t, (x, y), textcoords="offset points", xytext=(dx, dy), ha=ha,
                    va="bottom" if dy > 0 else "top", fontsize=9, color=INK)
    if logx:
        ax.set_xscale("log")
    if invert_x:
        ax.invert_xaxis()
    ax.grid(color=GRID, linewidth=0.8, zorder=0)
    ax.set_title(title, loc="left", fontsize=10.5, color=INK, fontweight="bold")
    ax.set_xlabel(xlabel)
    ax.set_ylabel(ylabel)
    lo, hi = min(ys), max(ys)
    ax.set_ylim(lo - (hi - lo) * 0.25, hi + (hi - lo) * 0.35)


FIG.mkdir(exist_ok=True)
fig, axes = plt.subplots(1, 3, figsize=(11, 3.6), dpi=200)
panel(axes[0], [r[3] for r in YOLOE11], [r[2] for r in YOLOE11], [f"11-{r[0]}\n{r[1]}M" for r in YOLOE11],
      "YOLOE-11 (텍스트 프롬프트)", "T4 TensorRT 속도 (FPS) → 빠름", "LVIS 제로샷 AP")
panel(axes[1], [r[3] for r in SAM21], [r[2] for r in SAM21], [f"{r[0]}\n{r[1]:g}M" for r in SAM21],
      "SAM 2.1", "A100 속도 (FPS) → 빠름", "SA-V test J&F",
      offsets={0: (0, -10, "center"), 1: (-8, -10, "right"), 3: (10, 9, "left")})
panel(axes[2], [r[1] for r in SIGLIP2], [r[2] for r in SIGLIP2], [r[0] for r in SIGLIP2],
      "SigLIP 2 (256px)", "비전 인코더 파라미터 (M, 로그)", "ImageNet 제로샷 정확도 (%)", logx=True)
axes[2].set_xticks([86, 303, 1000], ["86M", "303M", "1B"])
axes[2].xaxis.set_minor_locator(matplotlib.ticker.NullLocator())
fig.text(0.01, 0.01, "공식 공개 수치. 패널마다 측정 하드웨어·지표가 달라 패널끼리 직접 비교하면 안 된다.", fontsize=8, color=INK2)
fig.tight_layout(rect=(0, 0.05, 1, 1))
fig.savefig(FIG / "size_tradeoff.png", facecolor="white")
plt.close(fig)

# ---------------------------------------------------------------------------
# 그림: nooki 데이터 순환 구조 (YOLO11s 메인 + YOLOE 보조 + 사용자 확인)
# ---------------------------------------------------------------------------
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch

STEPS = [
    ("①  평소 탐색", "YOLO11s\n학습한 물품 검출", "로봇 · Hailo-8L"),
    ("②  목록 밖 요청", "YOLOE 텍스트 모드\n\"안경 찾아줘\"", "요청 시에만 실행"),
    ("③  사용자 확인", "대시보드 사진에\n맞아요 / 아니에요", "라벨 검증 ★"),
    ("④  데이터 축적", "확인된 사진 + 박스\n(필요 시 SAM 2로 보정)", "PC"),
    ("⑤  주기적 재학습", "요청이 쌓인 물품을\nYOLO11s 클래스에 추가", "기존 데이터와 함께"),
]
fig, ax = plt.subplots(figsize=(11, 3.3), dpi=200)
ax.set_xlim(0, 100)
ax.set_ylim(0, 34)
ax.axis("off")
bw, bh, gap, y0 = 17, 17, 3.25, 12
for i, (head, body, foot) in enumerate(STEPS):
    x = 1 + i * (bw + gap)
    key = i == 2
    ax.add_patch(FancyBboxPatch((x, y0), bw, bh, boxstyle="round,pad=0,rounding_size=1.2",
                                facecolor="#DEEAF6" if key else "white", edgecolor=SERIES, linewidth=2.2 if key else 1.2))
    ax.text(x + bw / 2, y0 + bh - 3, head, ha="center", va="center", fontsize=10, fontweight="bold", color=INK)
    ax.text(x + bw / 2, y0 + bh / 2 - 1, body, ha="center", va="center", fontsize=8.6, color=INK, linespacing=1.4)
    ax.text(x + bw / 2, y0 + 2.2, foot, ha="center", va="center", fontsize=7.8, color=INK2)
    if i < len(STEPS) - 1:
        ax.add_patch(FancyArrowPatch((x + bw + 0.4, y0 + bh / 2), (x + bw + gap - 0.4, y0 + bh / 2),
                                     arrowstyle="-|>", mutation_scale=12, color=INK2, linewidth=1.2))
# ⑤ → ① 되돌아가는 화살표
x_last = 1 + 4 * (bw + gap) + bw / 2
ax.add_patch(FancyArrowPatch((x_last, y0 - 0.5), (1 + bw / 2, y0 - 0.5), connectionstyle="bar,fraction=-0.09",
                             arrowstyle="-|>", mutation_scale=12, color=SERIES, linewidth=1.6))
ax.text(50, 6.2, "Hailo용 변환 후 재배포 → 다음부터는 YOLO11s가 직접 찾음", ha="center", va="center", fontsize=8.8,
        color=SERIES, bbox=dict(facecolor="white", edgecolor="none", pad=2))
ax.text(1, 32.5, "②는 YOLO11s 목록에 없는 물품을 요청받았을 때만 실행한다. 평소에는 ①만 돈다.",
        fontsize=8.3, color=INK2, va="center")
fig.tight_layout()
fig.savefig(FIG / "data_loop.png", facecolor="white")
plt.close(fig)

# ---------------------------------------------------------------------------
# 문서 서식
# ---------------------------------------------------------------------------
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2))


def set_font(style_or_run, size=None, bold=None, color=None):
    f = style_or_run.font
    f.name = FONT
    rpr = (style_or_run.element if hasattr(style_or_run, "element") else style_or_run._element).get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(k), FONT)
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color:
        f.color.rgb = RGBColor.from_string(color)


normal = doc.styles["Normal"]
set_font(normal, 10)
normal.paragraph_format.line_spacing = 1.25
normal.paragraph_format.space_after = Pt(7)

for lvl, size, before, after, color in [(1, 20, 0, 6, "1F3864"), (2, 14, 22, 8, "1F4E79"), (3, 11.5, 14, 5, "2E75B6")]:
    st = doc.styles[f"Heading {lvl}"]
    set_font(st, size, True, color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
    if lvl == 2:   # H2 아래 구분선
        ppr = st.element.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bot = OxmlElement("w:bottom")
        for k, v in {"w:val": "single", "w:sz": "6", "w:space": "4", "w:color": color}.items():
            bot.set(qn(k), v)
        bdr.append(bot)
        ppr.append(bdr)

for name in ("List Bullet", "List Number"):
    set_font(doc.styles[name], 10)


def H(text, lvl=2):
    doc.add_heading(text, lvl)


def P(text, size=None, color=None, italic=False, align=None):
    """**굵게** 표기를 지원하는 문단"""
    p = doc.add_paragraph()
    for i, part in enumerate(text.split("**")):
        if part:
            r = p.add_run(part)
            set_font(r, size, bold=(i % 2 == 1), color=color)
            r.italic = italic
    if align:
        p.alignment = align
    return p


def BL(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        for i, part in enumerate(it.split("**")):
            if part:
                set_font(p.add_run(part), bold=(i % 2 == 1))
        p.paragraph_format.space_after = Pt(3)


def shade(cell, hex_):
    tcpr = cell._element.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), hex_)
    tcpr.append(s)


def T(header, rows, widths=None, first_col_bold=True, center_cols=()):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    # 연회색 테두리
    tblpr = t._element.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        for k, v in {"w:val": "single", "w:sz": "4", "w:space": "0", "w:color": "BFBFBF"}.items():
            e.set(qn(k), v)
        borders.append(e)
    tblpr.append(borders)

    def fill(cell, text, bold=False, center=False):
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        for i, part in enumerate(str(text).split("**")):
            if part:
                set_font(p.add_run(part), 9, bold=bold or (i % 2 == 1))
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        fill(c, h, bold=True, center=True)
        shade(c, "DEEAF6")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            fill(cells[i], v, bold=(first_col_bold and i == 0), center=(i in center_cols))
    # 페이지를 넘어가면 머리글 행 반복, 행은 페이지 사이에서 쪼개지지 않게
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            trpr.append(OxmlElement("w:tblHeader"))
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def FIGURE(path, caption, width=16.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width))
    P(caption, size=8.5, color="7F7F7F", align=WD_ALIGN_PARAGRAPH.CENTER)


def NOTE(text):
    """회색 배경 한 칸 표로 만든 참고 상자"""
    t = doc.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]
    shade(c, "F2F2F2")
    c.text = ""
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for i, part in enumerate(text.split("**")):
        if part:
            set_font(p.add_run(part), 9, bold=(i % 2 == 1))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ---------------------------------------------------------------------------
# 본문
# ---------------------------------------------------------------------------
H("비전 AI 모델 5종 분석 리포트", 1)
P("YOLOE · SAM 2 · SigLIP 2 · EasyOCR · MediaPipe Hands — 어떤 상황에 어떤 모델이 맞는가", size=11, color="2E75B6")
P("작성일 2026-10-12 · 작성자 유정완", size=9, color="7F7F7F")

# 1 ---------------------------------------------------------------------------
H("1. 결론 요약")
P("다섯 모델은 서로 경쟁하는 모델이 아니라 **서로 다른 질문에 답하는 도구**다. "
  "YOLOE는 \"무엇이 어디에 있나\", SAM 2는 \"정확히 어디까지가 그 물체인가\", SigLIP 2는 \"이 사진과 이 말이 얼마나 비슷한가\", "
  "EasyOCR은 \"무슨 글자가 쓰여 있나\", MediaPipe Hands는 \"손이 어떤 모양인가\"에 답한다. "
  "그래서 \"어느 모델이 가장 좋은가\"보다 \"지금 풀려는 문제가 어떤 질문인가\"를 먼저 정해야 한다.")
P("공식 제원으로 보면 다섯 모델은 **어디서 돌리도록 설계됐는지**도 다르다.")
BL([
    "**MediaPipe Hands**는 처음부터 모바일 기기용으로 설계됐다. 입력이 192·224px로 작고, 공식 수치도 스마트폰(Pixel 6) CPU 기준 17.12ms다.",
    "**YOLOE**는 YOLO 구조를 그대로 써서 엣지 GPU·NPU용으로 만들어졌다. 공식 수치가 T4 GPU와 iPhone 12 기준이다.",
    "**SAM 2.1**은 가장 작은 tiny(38.9M)도 공식 속도가 A100 GPU 기준이다. 서버·PC용에 가깝다.",
    "**SigLIP 2**는 가장 작은 B/16도 다국어 어휘 때문에 텍스트 쪽이 커서 전체가 약 0.4B(4억) 파라미터다.",
    "**EasyOCR**은 공식 속도 수치가 없다. 범용 OCR 라이브러리로 GPU·CPU 모두 지원한다.",
])
P("nooki(실내 분실물 검출 로봇, 라즈베리 파이 5 + Hailo-8L)에 대한 결론은 다음과 같다.")
NOTE("**nooki 물품 검출의 결론: MVP는 학습한 YOLO11s, YOLOE는 확장 단계.** "
     "\"YOLOE가 nooki에 가장 적합하다\"는 뜻이 아니다. YOLOE는 다섯 모델 중 확장 기능(학습 안 한 물품을 텍스트로 찾기)에 가장 적합할 뿐이고, "
     "MVP의 물품 검출에는 nooki 물품으로 학습한 YOLO11s가 더 적합하다.")
T(["판단 기준", "MVP — 학습한 YOLO11s", "확장 — YOLOE-11-S 텍스트 모드"], [
    ["찾을 대상", "에어팟·반지·귀걸이·열쇠·리모컨 등 **정해진 목록** (PRD MVP 범위)", "학습하지 않은 **임의 물품** (PRD 비범위 · 추후 확장)"],
    ["정확도", "대상이 고정이면 그 물품으로 학습한 모델이 유리. YOLOE도 COCO로 다시 학습하면 AP 46.2로 일반 YOLO 수준이 됨",
     "학습 없이 찾는 제로샷 AP 27.5 (LVIS 1,203개 클래스) — 고정 목록에서는 학습 모델보다 불리"],
    ["작은 물체", "귀걸이·반지처럼 작은 물체는 실제 바닥 환경 데이터로 학습하는 쪽이 유리", "일반적인 이름 기반 탐지라 작은 물체·특수 환경에 대한 보장 없음"],
    ["배포", "Hailo-8L 공식 모델 목록에 있음", "Hailo-8L 가속 지원 **확인 필요**"],
    ["입력", "물품 목록에서 선택", "한국어 음성 → **영어 물품 이름 변환 단계 필요** (텍스트 인코더가 영어 중심)"],
    ["라이선스", "AGPL-3.0 (Ultralytics)", "AGPL-3.0 (원 저장소·Ultralytics) — 선택 기준이 되지 않음"],
], widths=[2.6, 7.0, 7.0])
P("두 모델은 하나만 고르는 관계가 아니다. **YOLOE를 nooki 물품 데이터로 학습하면 구조가 YOLO11과 같아진다**(공식 README: 전이 시 \"exactly the same\"). "
  "그래서 진행 순서는 ① MVP는 YOLO11s를 학습·배포하고, ② 확장 단계에서 YOLOE 텍스트 모드를 붙이되, ③ 붙이기 전에 Hailo-8L 변환 가능 여부를 먼저 검증하는 것이다.")
P("확장 단계에서 YOLOE는 **목록 밖 물품 요청을 처리하면서 새 물품의 학습 데이터를 모으는 역할**도 한다. "
  "사용자가 검출 결과를 \"맞아요\"로 확인한 사진을 모아 YOLO11s를 주기적으로 다시 학습하면, 자주 찾는 물품이 YOLO11s의 기본 목록으로 들어온다. "
  "단, 사용자 집 사진을 모으는 것은 PRD의 개인정보 원칙과 부딪히므로 데이터 수집 방식을 먼저 정해야 한다. (9-1)")
P("나머지 모델의 쓰임은 다음과 같다.")
BL([
    "**SigLIP 2는 2차 확인용으로 적합하다.** 다국어 학습이라 한국어 설명을 그대로 쓸 수 있고, "
    "설명 문장의 임베딩은 미리 계산해 둘 수 있어 로봇에는 비전 인코더(B/16 기준 86M)만 올리면 된다.",
    "**SAM 2는 로봇이 아니라 PC에서 데이터 라벨링 도구로 쓰는 것이 맞다.** 클릭으로 정밀 마스크를 만들어 학습 데이터 제작 시간을 줄인다.",
    "**MediaPipe Hands는 \"손가락 개수로 물품 고르기·탐색 시작\" 같은 비접촉 명령의 후보다.** 모바일 CPU용 설계라 Hailo-8L을 YOLO에 온전히 쓰면서 CPU에 둘 수 있다.",
    "**EasyOCR은 nooki의 핵심 기능과 거리가 멀다.** 글자가 있는 물품의 보조 확인 정도로 우선순위가 낮다.",
])
NOTE("**라이선스를 먼저 확인할 것.** YOLOE는 AGPL-3.0이다(원 저장소와 Ultralytics 패키지 모두). 제품으로 배포하려면 소스 공개 의무가 생기거나 "
     "Ultralytics 상용 라이선스가 필요하다. YOLOE 텍스트 인코더(MobileCLIP)의 가중치는 별도 약관(Apple ML Research Model TOU)이다. "
     "SAM 2·SigLIP 2·EasyOCR·MediaPipe는 Apache-2.0이다. (10장)")

# 2 ---------------------------------------------------------------------------
H("2. 다섯 모델 한눈에 보기")
T(["모델", "답하는 질문", "입력", "출력", "만든 곳 · 발표"], [
    ["YOLOE", "무엇이 어디에 있나?", "이미지 + (물체 이름 / 예시 박스 / 없음)", "클래스 이름 · 박스 · 마스크", "칭화대 · 2025"],
    ["SAM 2", "정확히 어디까지가 그 물체인가?", "이미지·영상 + 점 / 박스 / 마스크", "마스크 (이름 없음), 영상에서는 추적", "Meta · 2024 (2.1 개정)"],
    ["SigLIP 2", "이 사진과 이 말이 얼마나 비슷한가?", "이미지 / 텍스트", "임베딩 벡터 → 코사인 유사도", "Google DeepMind · 2025"],
    ["EasyOCR", "무슨 글자가 쓰여 있나?", "이미지", "글자 박스 · 텍스트 · 신뢰도", "JaidedAI (CRAFT·CRNN 기반)"],
    ["MediaPipe Hands", "손이 어떤 모양인가?", "이미지 · 영상 · 실시간 스트림", "손마다 관절 21개 (이미지·3D 좌표) · 왼손/오른손", "Google"],
], widths=[2.6, 3.6, 3.8, 3.8, 2.8])
P("다섯 모델을 **\"무엇(what)\"과 \"어디(where)\"** 두 축으로 보면 관계가 분명해진다.")
BL([
    "YOLOE는 무엇과 어디를 **한 번에** 답한다. 대신 경계는 상대적으로 거칠고, 이름은 텍스트 인코더가 이해하는 말(주로 영어 명사)로 줘야 한다.",
    "SAM 2는 **어디만** 아주 정밀하게 답한다. 무엇인지는 모르므로 탐지기와 짝을 이룬다.",
    "SigLIP 2는 **무엇만** 답한다. 이미지 전체를 벡터 하나로 요약하므로 위치는 모르지만, 정해진 클래스 목록 없이 자유 문장과 비교한다.",
    "EasyOCR과 MediaPipe는 **글자**와 **손**에 특화된 전문 모델이다. 대상이 좁은 대신 텍스트·관절 좌표처럼 바로 쓸 수 있는 구조화된 출력을 준다.",
])

# 3 ---------------------------------------------------------------------------
H("3. 공식 제원 비교")
P("각 모델의 공식 문서·논문에 공개된 값이다. **속도는 모델마다 측정 하드웨어가 달라서 서로 직접 비교할 수 없다.** "
  "같은 모델 안에서 크기를 고를 때의 참고값으로 읽어야 한다.")
T(["모델", "크기 구성", "대표 크기 파라미터", "입력 해상도", "공개 성능 (대표 크기)", "공개 속도와 측정 환경"], [
    ["YOLOE", "YOLOE-26: n/s/m/l/x\nYOLOE-11: s/m/l\nYOLOE-v8: s/m/l", "YOLOE-11-S\n10M (텍스트) / 12M (시각)", "640px",
     "LVIS 제로샷 AP 27.5 (텍스트) · 26.3 (시각)\n프롬프트 프리 20.6", "301.2 FPS (T4, TensorRT)\n73.3 FPS (iPhone 12, CoreML)"],
    ["SAM 2.1", "tiny / small / base+ / large", "tiny 38.9M\n(large 224.4M)", "1024px", "SA-V test J&F 76.5\n(large 79.5)", "91.2 FPS (A100)\n(large 39.5 FPS)"],
    ["SigLIP 2", "B / L / So400m / g\n(+ NaFlex 가변 해상도)", "B/16: 비전 86M\n전체 약 0.4B", "224~512px", "ImageNet 제로샷 79.1% (B/16, 256px)\n(g/16 256px 84.5%)", "공개 수치 없음"],
    ["EasyOCR", "언어별 인식 모델 + 공통 검출 모델", "공개 안 됨", "가변 (원본 크기 기반)", "80개 이상 언어 지원", "공개 수치 없음 (GPU·CPU 모두 지원)"],
    ["MediaPipe Hands", "단일 번들 (손바닥 검출 + 랜드마크)", "공개 안 됨\n(float16, 번들 약 7.8MB)", "192px (검출)\n224px (랜드마크)", "학습: 실제 이미지 약 3만 장 + 합성 손 모델", "17.12ms (Pixel 6 CPU)\n12.27ms (Pixel 6 GPU)"],
], widths=[2.3, 3.2, 2.9, 2.0, 3.4, 3.2])
P("이 표에서 읽어야 할 것은 세 가지다.")
BL([
    "**측정 하드웨어가 곧 설계 목표다.** MediaPipe는 스마트폰, YOLOE는 엣지 GPU와 iPhone, SAM 2.1은 데이터센터 GPU(A100)로 공식 수치를 냈다. "
    "어디서 돌릴지를 정하면 후보가 자연스럽게 좁혀진다.",
    "**SigLIP 2의 \"86M\"은 비전 인코더만의 크기다.** 다국어용 Gemma 토크나이저(어휘 256k)를 써서, 단어 임베딩 표만 약 256,000 × 768 ≈ 2억 개 파라미터다. "
    "그래서 B/16 전체가 약 0.4B가 된다. 텍스트 임베딩을 미리 계산해 두면 실행 시에는 비전 쪽만 있으면 된다.",
    "**제로샷 AP는 낮아 보여도 의미가 다르다.** YOLOE의 LVIS AP 27.5는 1,203개 클래스를 학습 없이 찾은 점수다. "
    "COCO 80개 클래스로 다시 학습하면 YOLOE-11-S도 AP 46.2로 일반 YOLO 수준이 된다.",
])

# 4 ---------------------------------------------------------------------------
H("4. 같은 모델 안에서 크기 고르기")
P("YOLOE, SAM 2.1, SigLIP 2는 여러 크기로 공개되어 있다. 크기를 키울 때 얻는 정확도와 잃는 속도를 공식 수치로 그리면 다음과 같다.")
FIGURE(FIG / "size_tradeoff.png", "그림 1. 모델 크기별 정확도와 속도. 점마다 크기 이름과 파라미터 수를 적었다. 패널끼리는 하드웨어·지표가 달라 비교하지 않는다.")
BL([
    "**YOLOE-11: S → M이 가장 효율적인 한 단계다.** 파라미터가 2배(10M → 21M)가 되면서 AP가 5.5 오르고(27.5 → 33.0) 속도는 약 44% 줄어든다(301 → 168 FPS). "
    "M → L은 AP 2.2를 더 얻는 데 그친다. 엣지 기기에서는 S로 시작하고, 정확도가 부족할 때 M을 검토하는 순서가 맞다.",
    "**SAM 2.1: tiny와 small은 정확도가 거의 같다(J&F 76.5 vs 76.6).** large는 J&F가 3.0 오르지만 파라미터가 약 5.8배이고 속도는 절반 이하다(91 → 40 FPS). "
    "라벨링 도구처럼 사람이 결과를 보고 고치는 용도라면 tiny로 충분하다.",
    "**SigLIP 2: 크기를 키울수록 효과가 줄어든다.** B → L(파라미터 3.5배)에서 3.4%p 오르지만, So400m → g(2.5배)에서는 1.1%p에 그친다. "
    "검색·2차 확인 용도라면 B나 L이 합리적이다.",
])
NOTE("**YOLOE-26**도 공개되어 있다(n~x 다섯 크기). 공식 문서 기준 YOLOE-26s는 10.7M 파라미터에 LVIS 텍스트 프롬프트 AP 30.8로, 비슷한 크기의 YOLOE-11-S(27.5)보다 높다. "
     "다만 텍스트 인코더가 mobileclip2_b(약 254MB)로 바뀌었고, 같은 기준의 속도표는 공개되지 않아 그림에서는 제외했다.")

# 5 ---------------------------------------------------------------------------
H("5. 능력 비교")
P("●는 그 모델의 주 기능, ◐는 가능하지만 제약이 있음, ○는 불가능을 뜻한다. 공식 문서에 적힌 기능과 모델 구조를 기준으로 했다.")
T(["능력", "YOLOE", "SAM 2", "SigLIP 2", "EasyOCR", "MediaPipe"], [
    ["물체 이름 알아내기", "●", "○", "◐ 후보 문장 중 고르기", "○", "○"],
    ["위치(박스) 찾기", "●", "◐ 프롬프트 필요", "○", "● 글자만", "● 손만"],
    ["정밀한 경계(마스크)", "◐ 인스턴스 분할", "●", "○", "○", "○"],
    ["학습 안 한 대상", "● 이름·예시로", "● 무엇이든", "● 자유 문장", "◐ 지원 언어만", "○ 손 전용"],
    ["프롬프트 언어", "◐ 영어 중심 (MobileCLIP)", "해당 없음", "● 다국어 (어휘 256k)", "● 80개 이상 언어", "해당 없음"],
    ["영상 처리", "◐ 프레임별 검출 + 추적기", "● 메모리 기반 추적", "○", "○", "● VIDEO·LIVE_STREAM 모드"],
    ["공식 타깃 하드웨어", "엣지 GPU · 모바일", "데이터센터 GPU", "GPU · TPU", "GPU · CPU", "모바일 CPU · GPU"],
    ["재학습·파인튜닝", "● YOLO 학습 도구 그대로", "◐ 가능하나 무거움", "◐ 가능", "◐ 인식 모델 학습 가능", "○ 공개 학습 도구 없음"],
    ["출력 형태", "박스 · 클래스 · 마스크", "마스크", "벡터", "텍스트", "관절 좌표"],
], widths=[3.4, 2.8, 2.6, 2.8, 2.5, 2.7], center_cols=(1, 2, 3, 4, 5))

# 6 ---------------------------------------------------------------------------
H("6. 모델별 분석")

H("6-1. YOLOE — 무엇이든 찾는 실시간 탐지", 3)
P("기존 YOLO는 학습한 클래스(COCO 80개 등)만 찾는다. YOLOE는 **찾을 대상을 실행할 때 알려줄 수 있는** 오픈 어휘 탐지·분할 모델이다. 알려주는 방법이 세 가지다.")
T(["모드", "핵심 기법", "특징"], [
    ["텍스트 프롬프트", "RepRTA — 텍스트 임베딩을 분류 헤드에 재파라미터화", "클래스를 정한 뒤에는 \"추론·전이 비용 0\"으로 일반 YOLO와 같은 구조가 됨"],
    ["시각 프롬프트", "SAVPE — 예시 박스를 의미 임베딩으로 변환", "이름 없이 예시 영역만으로 비슷한 물체를 찾음. 파라미터가 약간 늘어남 (11-S: 10M → 12M)"],
    ["프롬프트 프리", "LRPC — 내장 어휘 4,585개와 영역을 대조", "지시 없이 보이는 물체에 이름을 붙임. AP는 낮아짐 (11-S: 20.6)"],
], widths=[2.8, 6.2, 7.6])
P("**제원에서 보이는 강점**")
BL([
    "같은 크기의 이전 오픈 어휘 모델보다 효율적이다. 논문 기준 YOLOE-v8-S는 YOLO-Worldv2-S보다 LVIS AP가 3.5 높으면서, 학습 비용은 1/3이고 추론은 1.4배 빠르다.",
    "COCO로 전이 학습할 때 구조가 YOLOv8·YOLO11과 \"완전히 같다\"(공식 README). 그래서 Ultralytics의 학습·내보내기 도구를 그대로 쓰고, 일반 YOLO 배포 경로를 따를 수 있다.",
    "탐지와 인스턴스 분할 마스크를 한 번에 낸다.",
])
P("**한계**")
BL([
    "YOLOE-11의 텍스트 인코더는 MobileCLIP(mobileclip_blt)이다. CLIP 계열 텍스트 인코더는 주로 영어 캡션으로 학습되므로, 한국어·긴 문장보다 **짧은 영어 명사**에서 잘 동작한다.",
    "제로샷 AP(11-S 27.5)는 학습한 전용 모델보다 낮다. 클래스가 고정되어 있고 데이터가 있다면 전용 학습 모델이 더 정확하다.",
    "새 클래스 이름을 받을 때마다 텍스트 인코더가 필요하다. 인코더 파일이 모델 본체보다 크다(YOLOE-26용 mobileclip2_b 약 254MB).",
    "라이선스가 AGPL-3.0이다.",
])
P("**적합한 상황과 이유**")
BL([
    "**찾을 대상이 자주 바뀌는 탐지** — 창고 재고, 분실물, 현장 점검 대상이 매번 다를 때. 클래스마다 데이터를 모아 재학습하는 비용이 사라진다.",
    "**학습 데이터 초벌 라벨링** — 프롬프트 프리·텍스트 모드로 자동 박스를 만들고 사람이 고친다.",
    "**말로 설명하기 어려운 부품 찾기** — 시각 프롬프트로 사진 한 장의 예시만 주면 된다.",
    "부적합: 클래스가 고정되어 있고 데이터가 충분하다면 같은 크기의 일반 YOLO를 학습하는 편이 정확도·단순함 모두 유리하다.",
])

H("6-2. SAM 2 — 가리키면 잘라주는 범용 분할", 3)
P("점·박스·마스크로 가리키면 그 물체의 정확한 경계를 돌려준다. 계층형 비전 인코더(Hiera)가 이미지를 한 번 인코딩하면, "
  "가벼운 마스크 디코더가 프롬프트마다 마스크를 만든다. 이미지를 \"프레임 1개짜리 영상\"으로 다루고, 영상에서는 **메모리 어텐션**이 이전 프레임의 마스크를 기억해 같은 물체를 계속 추적한다.")
P("**제원에서 보이는 강점**")
BL([
    "학습 없이 어떤 물체든 분할한다. 특정 도메인 데이터가 없어도 쓸 수 있다.",
    "인코더와 디코더가 분리되어 있어, 같은 이미지에 프롬프트를 여러 번 바꿔 넣는 **대화형 작업**에 구조적으로 유리하다. 인코딩 비용은 이미지당 한 번만 든다.",
    "음성 점(\"여기는 물체가 아님\")으로 결과를 사람이 즉시 고칠 수 있다.",
    "Apache-2.0으로 상업 이용이 자유롭다.",
])
P("**한계**")
BL([
    "물체 이름을 모른다. 무엇인지는 다른 모델이 알려줘야 한다.",
    "가장 작은 tiny(38.9M)도 공식 속도가 A100 기준이고, 입력이 1024px로 크다. 엣지 기기에서 매 프레임 새 이미지를 처리하기에는 무겁다.",
    "프롬프트 없이 전부 분할하는 모드는 이미지 위에 격자로 수많은 점을 찍어 디코더를 반복 실행하므로 대화형 사용보다 훨씬 무겁다.",
])
P("**적합한 상황과 이유**")
BL([
    "**라벨링 도구** — 클릭 몇 번으로 픽셀 단위 마스크·박스를 만든다. tiny와 small의 정확도 차이가 거의 없어(그림 1) PC에서 tiny로 충분하다.",
    "**탐지기 뒤의 정밀화 단계** — 로봇 파지처럼 물체의 정확한 윤곽·면적·중심이 필요할 때, 탐지 박스를 SAM 2에 넣어 마스크로 바꾼다.",
    "**영상 속 특정 물체 하나 따라가기** — 첫 프레임에 점 하나만 찍으면 메모리로 이후 프레임을 추적한다.",
])

H("6-3. SigLIP 2 — 이미지와 텍스트를 같은 공간에", 3)
P("이미지와 텍스트를 같은 벡터 공간에 올려 코사인 유사도로 \"얼마나 비슷한가\"를 계산한다. CLIP은 softmax로 \"N개 문장 중 정답 하나\"를 고르도록 학습하지만, "
  "SigLIP은 **sigmoid로 (이미지, 문장) 쌍마다 독립적으로** 맞다/아니다를 판단하도록 학습한다. SigLIP 2는 여기에 캡션 생성 디코더, 자기 증류, 마스크 예측 손실을 더해 위치 파악과 세밀한 특징을 개선했다.")
P("**제원에서 보이는 강점**")
BL([
    "다국어 학습이다. Gemma 토크나이저(어휘 256k)를 쓰고 학습 데이터의 10%가 비영어 웹 페이지라 **한국어 설명을 번역 없이** 쓸 수 있는 후보다.",
    "모든 크기에서 이전 SigLIP보다 제로샷 분류·검색 성능이 높다(논문). 예: B/16 256px ImageNet 제로샷 79.1%.",
    "크기(B~g)와 해상도(224~512px)를 고를 수 있고, NaFlex 변형은 원본 종횡비를 유지해 문서·OCR성 이미지에 유리하다.",
    "임베딩을 미리 저장해 두면 검색은 벡터 내적 한 번이라, 수천 장도 즉시 비교한다. Apache-2.0이다.",
])
P("**한계**")
BL([
    "위치를 모른다. 이미지 전체를 하나의 벡터로 요약한다.",
    "텍스트 쪽이 크다. B/16도 다국어 어휘 때문에 전체 약 0.4B 파라미터다.",
    "CLIP 계열은 \"무엇이 있나\"에는 강하지만 \"버스 왼쪽의 남자\" 같은 위치 관계·개수 이해는 약한 편이다.",
    "sigmoid 확률은 엄격하게 나오는 경향이 있어, 절대값보다 후보 간 순위로 판단하는 것이 안전하다.",
])
P("**적합한 상황과 이유**")
BL([
    "**탐지 결과의 2차 확인** — 탐지기가 찾은 후보를 잘라 사용자의 설명과 비교한다. 위치는 탐지기가, 의미는 SigLIP 2가 맡는 분업이다.",
    "**사진 검색·중복 제거** — \"검은 지갑 사진\" 같은 텍스트 검색, 비슷한 사진 묶기.",
    "**제로샷 분류로 사전 검토** — 데이터를 모으기 전에 후보 클래스 문장만으로 분류 가능성을 빠르게 가늠한다.",
])

H("6-4. EasyOCR — 이미지 속 글자 읽기", 3)
P("글자 위치를 찾는 **검출(CRAFT)**과 잘라낸 영역의 글자를 읽는 **인식(ResNet·VGG 특징 추출 + LSTM + CTC)** 두 단계로 동작한다. "
  "CRAFT는 픽셀마다 \"글자 중심\"과 \"글자 사이 연결부\"일 확률을 히트맵으로 예측해 박스를 만든다. 인식 모델은 언어 그룹별로 따로 있다.")
P("**제원에서 보이는 강점**")
BL([
    "80개 이상 언어를 바로 쓸 수 있고 한국어를 지원한다. 영어는 모든 언어와 함께 쓸 수 있다.",
    "pip 설치 한 번으로 쓰는 완성형 라이브러리다. GPU·CPU 모두 지원하고 Apache-2.0이다.",
    "검출 박스가 네 꼭짓점 다각형이라 기울어진 글자도 표현한다.",
])
P("**한계**")
BL([
    "모든 언어를 동시에 쓸 수는 없다. 문자 체계가 다른 언어(예: 한국어와 일본어)는 한 번에 함께 지정하지 못한다.",
    "공식 속도 수치가 없고, 마지막 릴리스가 2024년 9월(1.7.2)로 개발 속도가 느린 편이다.",
    "검출 → 인식 2단계 구조라, 검출이 놓친 글자는 인식도 못 한다. 작은 글씨·기울어짐·흐림에서 검출 품질이 정확도를 좌우한다.",
    "숫자 0과 알파벳 O, 1과 l처럼 모양이 비슷한 글자를 혼동하는 것은 OCR 전반의 한계라 형식 검사 같은 후처리가 필요하다.",
])
P("**적합한 상황과 이유**")
BL([
    "**라벨·표지판·계기판 읽기** — 정해진 위치의 글자를 가끔 읽고, 결과를 형식(날짜, 번호 자릿수)으로 검증할 수 있을 때.",
    "**문서·영수증 디지털화** — 카메라 각도와 조명을 통제할 수 있는 환경.",
    "부적합: 움직이는 카메라로 매 프레임 글자를 읽는 용도. 글자 영역을 다른 모델로 먼저 좁히는 구성이 필요하다.",
])

H("6-5. MediaPipe Hands — 손 관절 21개 추적", 3)
P("**손바닥 검출 모델(192px)**이 손을 찾고, **랜드마크 모델(224px)**이 손 영역에서 관절 21개의 좌표를 회귀한다. "
  "손가락보다 모양 변화가 적은 손바닥을 먼저 찾는 설계라 검출이 안정적이다. 영상 모드에서는 이전 프레임의 손 위치를 이용해 검출 단계를 건너뛴다.")
P("**제원에서 보이는 강점**")
BL([
    "모바일용 설계다. float16 모델에 공식 수치가 Pixel 6 CPU 17.12ms, GPU 12.27ms다.",
    "출력이 풍부하다. 이미지 좌표, 미터 단위 3D 월드 좌표, 왼손/오른손 판정을 함께 준다. 월드 좌표는 원근 왜곡이 없어 손가락 각도·거리 계산에 적합하다.",
    "IMAGE·VIDEO·LIVE_STREAM 세 실행 모드와 Android·iOS·Python·Web을 모두 지원한다. Apache-2.0이다.",
])
P("**한계**")
BL([
    "손 전용이다. 다른 대상에는 쓸 수 없다.",
    "제스처의 \"의미\"(몇 개를 폈나, 어떤 신호인가)는 주지 않는다. 관절 좌표 위에 규칙이나 분류기를 직접 얹어야 한다.",
    "학습 데이터가 실제 이미지 약 3만 장 + 합성 데이터다. 손이 가려지면 보이지 않는 관절을 추정하므로 틀릴 수 있다.",
    "파라미터 수와 재학습 도구는 공개되어 있지 않다. 1.x 버전부터 예전 mp.solutions API가 없어지고 Tasks API로 바뀌었다.",
])
P("**적합한 상황과 이유**")
BL([
    "**비접촉 인터페이스** — 손가락 개수·제스처로 기기에 명령. 모바일 CPU용 설계라 별도 가속기 없이 돌릴 수 있는 후보다.",
    "**작업 동작 분석** — 손 위치·자세를 기록해 조립 동작, 수어 등을 분석.",
])

# 7 ---------------------------------------------------------------------------
H("7. 상황별 모델 선택 가이드")
T(["이런 상황이라면", "추천", "이유"], [
    ["정해진 물체를 엣지에서 계속 찾는다", "일반 YOLO (클래스 학습)", "오픈 어휘가 필요 없으면 같은 크기에서 가장 정확하고 단순함"],
    ["찾을 물체가 자주 바뀐다", "YOLOE 텍스트 모드 (11-S부터)", "재학습 없이 이름만 바꿔 넣음. 클래스 고정 후 YOLO와 같은 구조"],
    ["이름은 모르지만 예시 사진은 있다", "YOLOE 시각 프롬프트", "박스 하나로 예시를 지정"],
    ["이미지에 뭐가 있는지 전부 알고 싶다", "YOLOE 프롬프트 프리", "내장 어휘 4,585개로 넓게 탐지 (정확도는 낮음)"],
    ["물체의 정확한 윤곽·면적이 필요하다", "탐지기 + SAM 2.1", "박스를 정밀 마스크로 바꿈"],
    ["학습용 마스크·박스를 빨리 만들어야 한다", "SAM 2.1 tiny (PC)", "인코더 한 번 + 가벼운 디코더 구조. tiny와 small 정확도가 거의 같음"],
    ["사용자의 자유로운 설명과 사진을 비교한다", "SigLIP 2 (B 또는 L)", "다국어, 자유 문장, 임베딩 재사용"],
    ["비슷한 사진을 찾거나 묶는다", "SigLIP 2 이미지 임베딩", "벡터 내적으로 대량 비교"],
    ["라벨·표지판 글자를 읽는다", "EasyOCR", "80개 이상 언어, 한국어 지원"],
    ["손짓으로 기기를 조작한다", "MediaPipe Hands", "모바일 CPU용 설계, 관절 좌표 출력"],
], widths=[5.2, 4.2, 7.2])
P("결정이 애매할 때는 다음 질문 순서로 좁히면 된다.")
BL([
    "출력이 무엇이어야 하나? — 이름(YOLOE·SigLIP 2), 위치(YOLOE), 경계(SAM 2), 텍스트(EasyOCR), 관절(MediaPipe)",
    "대상이 고정인가, 바뀌는가? — 고정이면 학습한 전용 모델, 바뀌면 오픈 어휘 모델",
    "어디서 돌리나? — 3장 표의 공식 측정 하드웨어가 그 모델의 설계 목표다. 모바일(MediaPipe), 엣지 GPU·NPU(YOLOE), 서버 GPU(SAM 2.1)",
    "같은 모델이면 어느 크기인가? — 그림 1에서 정확도가 꺾이는 지점(YOLOE S→M, SAM tiny≈small, SigLIP B~L)을 기준으로",
], style="List Number")

# 8 ---------------------------------------------------------------------------
H("8. 조합해서 쓰기")
P("한 모델로 끝나는 문제는 드물다. 각 모델의 약점을 다른 모델이 메우는 조합이 실전에서 자주 쓰인다.")
T(["조합", "흐름", "메우는 약점"], [
    ["YOLOE → SAM 2", "YOLOE가 박스와 이름 → SAM 2가 박스를 정밀 마스크로", "YOLOE의 거친 경계, SAM의 이름 없음"],
    ["YOLO/YOLOE → SigLIP 2", "탐지 후보를 잘라 사용자 설명과 유사도 비교", "탐지기의 오검출, 색·모양 같은 세부 조건"],
    ["음성 → LLM → YOLOE", "\"안경 찾아줘\" → 영어 명사 [\"glasses\"] → 클래스 설정", "YOLOE 텍스트 인코더의 한국어·문장 약점"],
    ["YOLOE → EasyOCR", "글자가 있을 물체(라벨, 표지판)만 잘라 OCR", "OCR 검출 범위를 좁혀 오검출·연산 감소"],
    ["MediaPipe → 기기 명령", "손가락 개수·제스처 → 메뉴 선택·시작/정지", "버튼·음성이 어려운 환경의 입력 수단"],
], widths=[3.6, 7.0, 6.0])

# 9 ---------------------------------------------------------------------------
H("9. nooki 적용 제안")
P("nooki PRD의 조건은 ① 모든 처리를 로봇 안(라즈베리 파이 5 + Hailo-8L, ROS 2)에서 하고, ② MVP는 학습한 물품 목록만 찾으며, "
  "③ \"텍스트로 아무 물건이나 찾기\"는 추후 확장이라는 것이다. 이 조건과 공식 제원으로 다섯 모델의 쓰임을 정리하면 다음과 같다. "
  "1장 결론대로 **물품 검출의 중심은 MVP의 학습한 YOLO11s이고, YOLOE는 확장 단계의 추가 기능**이다.")
T(["단계", "모델", "역할", "실행 위치", "우선순위"], [
    ["데이터 준비", "SAM 2.1 tiny", "에어팟·반지·귀걸이 사진에 클릭으로 박스·마스크 라벨링", "PC (오프라인)", "높음"],
    ["데이터 준비", "YOLOE 프롬프트 프리·텍스트", "수집한 바닥 영상에 자동 초벌 라벨", "PC (오프라인)", "높음"],
    ["탐색 시작", "MediaPipe Hands", "손가락 개수로 찾을 물품 번호 선택, 손바닥으로 시작", "로봇 CPU", "중간 (음성·대시보드 보조)"],
    ["물품 검출", "학습한 YOLO11s", "MVP의 지정 물품 검출 (기존 보고서 추천안)", "로봇 Hailo-8L", "MVP 핵심"],
    ["물품 검출 (확장)", "YOLOE-11-S 텍스트 모드", "학습 안 한 물품을 영어 이름으로 검출", "로봇 (가속 지원 확인 필요)", "확장 1순위"],
    ["검출 확인", "SigLIP 2 B/16", "검출 후보 사진과 \"검은색 이어폰 케이스\" 같은 설명 비교", "로봇 또는 대시보드", "확장 2순위"],
    ["보조 확인", "EasyOCR", "글자가 있는 물품(리모컨 등) 확인", "대시보드", "낮음"],
], widths=[2.4, 3.4, 5.4, 3.0, 2.4])
P("**제원에서 나오는 설계 포인트**")
BL([
    "**YOLOE는 \"텍스트 → 임베딩\"만 로봇 밖으로 뺄 수 있다.** 텍스트 인코더는 새 이름을 받을 때만 필요하고, 클래스를 정한 뒤의 탐지 모델은 일반 YOLO11과 같은 구조다. "
    "대시보드에서 이름을 임베딩으로 바꿔 로봇에 보내면, 로봇에는 큰 텍스트 인코더를 올리지 않아도 된다.",
    "**SigLIP 2도 같은 방식이 된다.** 물품 설명 문장의 임베딩을 미리 계산해 두면 로봇에는 비전 인코더(B/16 기준 86M)만 있으면 된다. 0.4B 전체를 올릴 필요가 없다.",
    "**MediaPipe는 모바일 CPU용 설계라 CPU에 두고, Hailo-8L은 YOLO에 온전히 쓰는 분담**이 자연스럽다.",
    "**SAM 2.1은 로봇에 올리지 않는다.** 가장 작은 tiny도 1024px 입력에 데이터센터 GPU 기준 수치라, 엣지 실시간용이 아니라 PC 라벨링용으로 쓴다.",
    "**Hailo-8L 가속 지원은 확인이 필요하다.** 일반 YOLO는 Hailo 공식 모델 목록에 있지만, YOLOE·SigLIP 2가 변환·가속되는지는 공식적으로 확인되지 않았다.",
    "**PRD의 \"영상을 외부로 보내지 않는다\" 원칙**상, SigLIP 2를 대시보드에서 돌린다면 전체 영상이 아니라 잘라낸 후보 이미지만 보내는 것이 허용되는지 정해야 한다.",
])

H("9-1. YOLO11s 메인 + YOLOE 보조로 학습 데이터 모으기", 3)
P("YOLO11s를 메인 검출기로 두고, YOLOE는 **목록에 없는 물품 요청을 처리하면서 새 물품의 학습 데이터를 모으는 역할**로 쓰는 구조다. "
  "모은 데이터로 YOLO11s를 주기적으로 다시 학습하면, 사용자가 자주 찾는 물품이 점점 YOLO11s의 기본 목록으로 들어온다.")
FIGURE(FIG / "data_loop.png", "그림 2. nooki 데이터 순환 구조. ③ 사용자 확인이 라벨을 검증하는 핵심 단계다.")
T(["단계", "하는 일", "실행 위치"], [
    ["① 평소 탐색", "YOLO11s가 학습한 물품(에어팟·반지·귀걸이·열쇠·리모컨)을 찾는다", "로봇 (Hailo-8L)"],
    ["② 목록 밖 요청", "\"안경 찾아줘\"처럼 목록에 없는 물품을 요청받으면 YOLOE 텍스트 모드로 대신 탐색한다", "요청 시에만 실행"],
    ["③ 사용자 확인", "대시보드에 뜬 검출 사진에 사용자가 \"맞아요 / 아니에요\"를 누른다", "대시보드"],
    ["④ 데이터 축적", "\"맞아요\"로 확인된 사진과 박스를 학습 데이터로 저장한다. 박스가 부정확하면 SAM 2로 다듬는다", "PC"],
    ["⑤ 주기적 재학습", "요청이 일정 횟수 넘게 쌓인 물품을 YOLO11s 클래스에 추가해 다시 학습·변환·배포한다", "PC → 로봇"],
], widths=[3.0, 10.4, 3.2])
P("**핵심은 ③ 사용자 확인이다.** YOLOE의 검출은 학습 없이 찾은 결과라 정확도가 낮고(YOLOE-11-S 제로샷 AP 27.5), 작은 물체에서는 더 불리하다. "
  "확인 없이 그대로 학습 데이터로 쓰면 YOLOE의 실수까지 YOLO11s가 배운다. 반대로 사용자가 \"맞아요\"를 누르는 순간 라벨이 검증되므로, "
  "따로 라벨링하는 사람 없이도 믿을 수 있는 데이터가 쌓인다.")
P("**주의할 점**")
T(["주의할 점", "이유", "대응"], [
    ["PRD 개인정보 원칙과 충돌", "PRD는 집 내부 영상을 외부로 보내지 않고, 영상의 클라우드 전송·저장을 비범위로 정했다. 사용자 집 사진을 모아 재학습하려면 이 원칙을 바꿔야 한다",
     "셋 중 하나를 정한다: ① 동의를 받고 확인된 잘라낸 사진만 전송 ② 데이터 수집은 팀이 꾸민 실험 공간에서만 ③ 로봇 안에서만 저장·재학습 (구현 난도 높음)"],
    ["YOLO11s와 YOLOE 동시 상시 실행 금지", "Hailo-8L 하나에 두 모델을 매 프레임 돌리면 처리 속도가 떨어진다", "YOLOE는 ② 목록 밖 요청이 있을 때만 실행하거나, 저장한 프레임을 PC에서 처리"],
    ["재학습 시 기존 데이터 함께 사용", "새 물품 데이터만으로 학습하면 기존 물품 검출 성능이 떨어질 수 있다", "기존 물품 데이터와 새 데이터를 합쳐서 학습하고, 기존 물품 정확도를 배포 전에 다시 검증"],
    ["재학습마다 변환·재배포 필요", "새 클래스를 추가할 때마다 Hailo용 변환(컴파일)과 로봇 배포를 다시 해야 한다", "\"요청 N회 이상 · 확인 데이터 M장 이상\" 같은 추가 기준을 정해 주기적으로 묶어서 진행"],
    ["② 단계의 실행 위치 미확정", "Hailo-8L에서 YOLOE를 가속할 수 있는지 확인되지 않았다", "먼저 변환 가능 여부를 검증. 안 되면 ②를 대시보드에서 처리하는 구조를 검토하되, 이 경우 첫 번째 행의 개인정보 문제와 함께 결정"],
], widths=[3.4, 6.4, 6.8])

# 10 ---------------------------------------------------------------------------
H("10. 라이선스")
T(["모델 · 구성 요소", "라이선스", "제품 적용 시 주의"], [
    ["YOLOE (THU-MIG 원 저장소)", "AGPL-3.0", "배포·서비스 시 소스 공개 의무"],
    ["Ultralytics 패키지 (YOLOE·SAM 2 실행)", "AGPL-3.0 또는 Ultralytics Enterprise", "상용 제품은 Enterprise 라이선스 검토"],
    ["MobileCLIP (YOLOE 텍스트 인코더)", "코드 MIT · 모델 Apple ML Research Model TOU · 데이터 CC-BY-NC-ND", "가중치 약관의 상업 이용 조건 확인 필요"],
    ["SAM 2 / 2.1 (Meta 원 저장소)", "Apache-2.0", "Meta 공식 코드로 쓰면 Apache-2.0만 적용"],
    ["SigLIP 2", "Apache-2.0", "자유롭게 사용, 고지 의무"],
    ["EasyOCR", "Apache-2.0", "자유롭게 사용, 고지 의무"],
    ["MediaPipe", "Apache-2.0", "자유롭게 사용, 고지 의무"],
], widths=[5, 5.4, 6.2])

# 11 ---------------------------------------------------------------------------
H("11. 용어 정리")
T(["용어", "뜻"], [
    ["오픈 어휘 (open-vocabulary)", "학습 때 정한 클래스 목록에 갇히지 않고, 실행할 때 준 이름·예시로 대상을 찾는 방식"],
    ["제로샷 (zero-shot)", "그 작업용 학습 데이터 없이 바로 수행하는 것"],
    ["LVIS AP", "1,203개 클래스 데이터셋 LVIS에서 잰 탐지 정확도. 희귀 클래스가 많아 COCO AP보다 낮게 나옴"],
    ["J&F", "영상 분할 정확도 지표. 영역 겹침(J)과 경계 정확도(F)의 평균"],
    ["재파라미터화", "학습 때 쓴 추가 구조를 추론 전에 기존 가중치에 합쳐서, 실행 시 비용을 없애는 기법"],
    ["임베딩", "이미지·텍스트를 고정 길이 숫자 벡터로 바꾼 것. 벡터 사이 각도가 의미의 유사성을 나타냄"],
    ["코사인 유사도", "두 벡터 사이 각도의 코사인. 1이면 같은 방향, 0이면 무관"],
    ["토크나이저 어휘", "텍스트를 쪼개는 단위(토큰)의 목록. 어휘가 클수록 여러 언어를 다루지만 임베딩 표가 커짐"],
    ["분할 마스크", "물체에 해당하는 픽셀을 1, 나머지를 0으로 표시한 이미지. 박스보다 정밀함"],
    ["메모리 어텐션", "SAM 2가 이전 프레임의 정보를 저장·참조해 영상에서 같은 물체를 추적하는 구조"],
    ["CTC", "글자 위치 라벨 없이 문자열 인식을 학습하게 해주는 방법. 중복·공백을 정리해 최종 문자열을 만듦"],
    ["랜드마크", "MediaPipe가 출력하는 관절 위치 점. 손은 21개"],
    ["AGPL-3.0", "네트워크 서비스로 제공해도 소스 공개 의무가 생기는 강한 카피레프트 라이선스"],
], widths=[4.2, 12.4])

# 출처 ---------------------------------------------------------------------------
H("출처")
BL([
    "Wang et al., \"YOLOE: Real-Time Seeing Anything\", arXiv:2503.07465, 2025.",
    "THU-MIG/yoloe 공식 저장소 README — 크기별 파라미터, LVIS AP, T4·iPhone 12 FPS, COCO 전이 결과 (github.com/THU-MIG/yoloe)",
    "Ultralytics YOLOE 문서 — YOLOE-26 제원, 내장 어휘 4,585개, 텍스트 인코더 (docs.ultralytics.com/models/yoloe)",
    "Ravi et al., \"SAM 2: Segment Anything in Images and Videos\", arXiv:2408.00714, 2024.",
    "facebookresearch/sam2 공식 저장소 — SAM 2.1 크기별 파라미터, A100 FPS, SA-V·MOSE·LVOS J&F, 라이선스 (github.com/facebookresearch/sam2)",
    "Tschannen et al., \"SigLIP 2: Multilingual Vision-Language Encoders with Improved Semantic Understanding, Localization, and Dense Features\", arXiv:2502.14786, 2025.",
    "Zhai et al., \"Sigmoid Loss for Language Image Pre-Training\" (SigLIP), arXiv:2303.15343, 2023.",
    "Hugging Face, \"SigLIP 2: A better multilingual vision language encoder\" 블로그 · google/siglip2-base-patch16-224 모델 카드",
    "JaidedAI/EasyOCR 공식 저장소 — 지원 언어, 구조, 릴리스, 라이선스 (github.com/JaidedAI/EasyOCR)",
    "Baek et al., \"Character Region Awareness for Text Detection\" (CRAFT), arXiv:1904.01941, 2019.",
    "Shi et al., \"An End-to-End Trainable Neural Network for Image-based Sequence Recognition\" (CRNN), arXiv:1507.05717, 2015.",
    "Google, MediaPipe Hand Landmarker 문서 — 모델 구성, 입력 크기, Pixel 6 지연 시간, 학습 데이터 (developers.google.com/edge/mediapipe/solutions/vision/hand_landmarker)",
    "Zhang et al., \"MediaPipe Hands: On-device Real-time Hand Tracking\", arXiv:2006.10214, 2020.",
    "apple/ml-mobileclip 저장소 라이선스 · ultralytics/ultralytics 저장소 라이선스",
])

doc.save(OUT)
print(f"저장: {OUT}")

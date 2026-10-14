"""Nooki 보고서 docx 재디자인: 데이터시트 방향. usage: restyle.py in.docx out.docx"""
import sys, copy
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

INK, SLATE, RULE, FILL = "1C2733", "5B6B7B", "CBD3DC", "F2F5F8"
POW, DATA = "D9622B", "2B6CB0"
FONT = "Malgun Gothic"
TEXT_W = 9638  # A4 - 20mm*2, dxa

src, dst = sys.argv[1], sys.argv[2]
d = Document(src)

# ---------- 페이지 ----------
sec = d.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin = sec.right_margin = Mm(20)
sec.top_margin, sec.bottom_margin = Mm(22), Mm(20)
sec.header_distance, sec.footer_distance = Mm(10), Mm(10)

# ---------- 스타일 ----------
def set_fonts(rpr_parent_el, name=FONT):
    rpr = rpr_parent_el
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)

styles_el = d.styles.element
rpr_default = styles_el.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
set_fonts(rpr_default)
lang = rpr_default.find(qn("w:lang"))
if lang is not None:
    lang.set(qn("w:eastAsia"), "ko-KR")
# 숫자·영문과 한글 사이에 Word가 자동으로 넣는 간격 끄기 ("19 종" → "19종")
ppr_default = styles_el.find(qn("w:docDefaults")).find(qn("w:pPrDefault")).find(qn("w:pPr"))
for tag in ("w:autoSpaceDE", "w:autoSpaceDN"):
    e = OxmlElement(tag); e.set(qn("w:val"), "0"); ppr_default.insert(0, e)

def style(name, size, bold=False, color=INK, before=0, after=6, line=1.1, keep=False):
    s = d.styles[name]
    s.font.name = FONT; s.font.size = Pt(size); s.font.bold = bold
    s.font.color.rgb = RGBColor.from_string(color)
    set_fonts(s.element.get_or_add_rPr())
    pf = s.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
    if keep: pf.keep_with_next = True
    return s

style("Normal", 10, after=7, line=1.15)
style("Heading 2", 15, bold=True, before=24, after=8, keep=True)
style("Heading 3", 11, bold=True, before=14, after=5, keep=True)
style("List Paragraph", 10, after=4, line=1.15)
try:
    cap = d.styles["Report Caption"]
except KeyError:
    cap = d.styles.add_style("Report Caption", WD_STYLE_TYPE.PARAGRAPH)
cap.base_style = d.styles["Normal"]
style("Report Caption", 8.5, color=SLATE, before=2, after=14)
cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ---------- 본문 문단 ----------
paras = list(d.paragraphs)
title_p, byline_p = paras[0], paras[1]
lead_text = paras[3].text.split(". ")[0] + "."

prev_img = False
for p in paras[2:]:
    has_img = bool(p._p.xpath(".//w:drawing"))
    if has_img:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        prev_img = True
        continue
    if prev_img:
        p.style = d.styles["Report Caption"]; prev_img = False
        continue
    if p.text.lstrip().startswith("☐"):
        ppr = p._p.pPr
        if ppr is not None and ppr.numPr is not None:
            ppr.remove(ppr.numPr)
        p.paragraph_format.left_indent = Mm(4)
    runs = [r for r in p.runs if r.text.strip()]
    if p.style.name == "Normal" and runs and all(r.bold for r in runs) and len(p.text) < 40:
        p.style = d.styles["Heading 3"]
        for r in p.runs: r.bold = None

# ---------- 표 ----------
def tc_prop(cell):
    return cell._tc.get_or_add_tcPr()

def shade(cell, hex_):
    tcpr = tc_prop(cell)
    for old in tcpr.findall(qn("w:shd")): tcpr.remove(old)
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_)
    tcpr.append(shd)

def cell_borders(cell, **edges):
    tcpr = tc_prop(cell)
    b = tcpr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders"); tcpr.append(b)
    for edge, (sz, color) in edges.items():
        e = b.find(qn("w:" + edge))
        if e is None:
            e = OxmlElement("w:" + edge); b.append(e)
        e.set(qn("w:val"), "single" if sz else "nil"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), color); e.set(qn("w:space"), "0")

def table_base(t, widths):
    tbl = t._tbl; tblpr = tbl.tblPr
    for tag in ("w:tblStyle", "w:tblBorders", "w:tblCellMar", "w:tblW", "w:jc"):
        for old in tblpr.findall(qn(tag)): tblpr.remove(old)
    w = OxmlElement("w:tblW"); w.set(qn("w:w"), str(sum(widths))); w.set(qn("w:type"), "dxa"); tblpr.append(w)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "right", "insideV"):
        e = OxmlElement("w:" + edge); e.set(qn("w:val"), "nil"); borders.append(e)
    for edge, sz, col in (("bottom", 8, INK), ("insideH", 4, RULE)):
        e = OxmlElement("w:" + edge); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), col); e.set(qn("w:space"), "0"); borders.append(e)
    tblpr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for edge, v in (("top", 55), ("bottom", 55), ("left", 85), ("right", 85)):
        e = OxmlElement("w:" + edge); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblpr.append(mar)
    grid = tbl.tblGrid
    for gc, wv in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(wv))
    for row in t.rows:
        trpr = row._tr.get_or_add_trPr()
        if trpr.find(qn("w:cantSplit")) is None: trpr.append(OxmlElement("w:cantSplit"))
        for c, wv in zip(row.cells, widths):
            c.width = wv * 635  # dxa → EMU
    hdr = t.rows[0]._tr.get_or_add_trPr()
    hdr.append(OxmlElement("w:tblHeader"))

def cell_text(cell, size, color=INK, bold=None, align=None):
    for p in cell.paragraphs:
        pf = p.paragraph_format
        pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.05
        if align is not None: p.alignment = align
        for r in p.runs:
            r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
            if bold is not None: r.bold = bold

def style_table(t, widths, size=8.5, zebra=True, align=None):
    table_base(t, widths)
    for ri, row in enumerate(t.rows):
        for ci, c in enumerate(row.cells):
            a = (align or {}).get(ci)
            if ri == 0:
                shade(c, INK); cell_text(c, size, "FFFFFF", True, a)
            else:
                if zebra and ri % 2 == 0: shade(c, FILL)
                cell_text(c, size, INK, None, a)

R, C = WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER
parts, rails, comm, pins, risks = d.tables

# 부품 총괄표: 계층별로 묶고 셀 병합
style_table(parts, [760, 1900, 560, 1250, 940, 1010, 1400, 1100, 718], size=8, zebra=False, align={2: C, 4: R, 5: R})
rows = parts.rows
layer_col = {"전원": POW, "컴퓨팅": DATA}
groups, start = [], 1
for i in range(2, len(rows) + 1):
    if i == len(rows) or rows[i].cells[0].text != rows[start].cells[0].text:
        groups.append((start, i - 1)); start = i
for a, b in groups:
    name = rows[a].cells[0].text
    for ri in range(a, b + 1):
        for c in rows[ri].cells:
            if ri == a and a > 1: cell_borders(c, top=(8, SLATE))
    for ri in range(a + 1, b + 1):
        for p in rows[ri].cells[0].paragraphs:
            for r in list(p.runs): r._r.getparent().remove(r._r)
    if b > a:
        merged = rows[a].cells[0].merge(rows[b].cells[0])
        for p in merged.paragraphs[1:]:
            p._p.getparent().remove(p._p)
    first = rows[a].cells[0]
    for r in first.paragraphs[0].runs:
        r.bold = True; r.font.color.rgb = RGBColor.from_string(layer_col.get(name, INK))

style_table(rails, [1500, 2900, 1050, 1050, 1838, 1300], align={2: R, 3: R, 5: R})
for row in rails.rows[1:]:
    c = row.cells[5]
    try:
        if int(c.text.strip().rstrip("%")) >= 80:
            for r in c.paragraphs[0].runs: r.bold = True; r.font.color.rgb = RGBColor.from_string(POW)
    except ValueError:
        pass
style_table(comm, [1900, 2350, 1450, 3938])
style_table(pins, [3300, 1900, 4438])
style_table(risks, [1750, 3300, 4588])
for t in (comm, pins, risks):
    for row in t.rows[1:]:
        for r in row.cells[0].paragraphs[0].runs: r.bold = True

# ---------- 쪽 나눔 규칙 ----------
# 1) 큰 그림이나 긴 표로 시작하는 절만 새 쪽에서 시작 (나머지는 앞 절에 이어서 흘려 빈 쪽을 줄인다)
NEW_PAGE = {"시스템 구성", "하드웨어 관계도", "부품 총괄표", "통신 구간과 핀 배정"}
for p in d.paragraphs:
    if p.style.name == "Heading 2" and p.text.strip() in NEW_PAGE:
        p.paragraph_format.page_break_before = True

# 2) 그림·표 바로 앞 문단은 그림·표와 같은 쪽에
def keep_prev_with(el):
    prev = el.getprevious()
    if prev is not None and prev.tag == qn("w:p"):
        from docx.text.paragraph import Paragraph
        Paragraph(prev, d._body).paragraph_format.keep_with_next = True
for p in d.paragraphs:
    if p._p.xpath(".//w:drawing"):
        keep_prev_with(p._p)
for t in d.tables:
    keep_prev_with(t._tbl)

# 3) 작은 표(12행 이하)는 한 쪽에 통째로, 긴 표는 머리+첫 3행을 함께
def keep_rows(t, upto):
    for row in t.rows[:upto]:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True
for t in d.tables:
    n = len(t.rows)
    keep_rows(t, n - 1 if n <= 12 else 4)

# 4) 다이어그램은 본문 폭에 가깝게(165mm) 키워 작은 글자도 읽히게
for shp in d.inline_shapes:
    ratio = shp.height / shp.width
    shp.width = Mm(165); shp.height = int(Mm(165) * ratio)

# 표 앞뒤 간격
for t in d.tables:
    nxt = t._tbl.getnext()
    if nxt is not None and nxt.tag == qn("w:p"):
        pPr = nxt.find(qn("w:pPr"))
        if pPr is None:
            pPr = OxmlElement("w:pPr"); nxt.insert(0, pPr)
        sp = OxmlElement("w:spacing"); sp.set(qn("w:before"), "200"); pPr.append(sp)

# ---------- 표지 ----------
anchor = title_p
def cover_par(text="", size=10, color=INK, bold=False, before=0, after=0, line=1.1):
    p = anchor.insert_paragraph_before()
    pf = p.paragraph_format; pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = line
    if text:
        r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.font.color.rgb = RGBColor.from_string(color)
    return p

cover_par("비전 기반 실내 분실물 검출 무인이동체", 10.5, SLATE, before=40, after=8)
cover_par("Nooki 하드웨어", 30, INK, True, after=0, line=1.0)
cover_par("부품 정리 보고서", 30, INK, True, after=14, line=1.0)
cover_par(lead_text, 11, SLATE, after=24, line=1.3)

spec = d.add_table(rows=2, cols=5)
vals = [("19종", "부품 수"), ("≈18W", "평균 전력"), ("≈43W", "피크 전력"), ("≈100분", "1회 충전 사용"), ("5.4cm", "차체 높이")]
table_base(spec, [1928, 1928, 1928, 1927, 1927])
tblpr = spec._tbl.tblPr
for e in tblpr.find(qn("w:tblBorders")):
    e.set(qn("w:val"), "nil")
spec.rows[0]._tr.trPr.remove(spec.rows[0]._tr.trPr.find(qn("w:tblHeader")))
for i, (v, l) in enumerate(vals):
    top, bot = spec.rows[0].cells[i], spec.rows[1].cells[i]
    cell_borders(top, top=(18, INK)); cell_borders(bot, bottom=(4, RULE))
    top.text = v; bot.text = l
    cell_text(top, 20, POW if i in (1, 2, 3) else INK, True)
    cell_text(bot, 8.5, SLATE)
    top.paragraphs[0].paragraph_format.space_before = Pt(8)
    bot.paragraphs[0].paragraph_format.space_after = Pt(8)
anchor._p.addprevious(spec._tbl)

leg = cover_par(before=10, after=26)
for txt, col in (("■ ", POW), ("주황", POW), (" 숫자와 선은 전원,  ", SLATE), ("■ ", DATA), ("파랑", DATA), (" 선은 데이터와 두 보드를 뜻한다.", SLATE)):
    r = leg.add_run(txt); r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string(col)

cover_par("목차", 11, INK, True, after=6)
toc = cover_par(after=0)
r = toc.add_run()
for tag, attr in (("w:fldChar", "begin"), (None, None), ("w:fldChar", "separate"), ("t", None), ("w:fldChar", "end")):
    if tag == "w:fldChar":
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), attr); r._r.append(e)
    elif tag is None:
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = ' TOC \\o "2-2" \\h \\z \\u '; r._r.append(it)
    else:
        t = OxmlElement("w:t"); t.text = "목차를 갱신하려면 F9"; r._r.append(t)

cover_par("작성  유정완 (2026년 10월 14일)", 9, SLATE, before=24, after=2)
last = cover_par("근거 자료  Nooki PRD(2026.10.06), Physical AI 하드웨어 설계 인터랙티브 강의", 9, SLATE)
last.add_run().add_break(WD_BREAK.PAGE)

for p in (title_p, byline_p):
    p._p.getparent().remove(p._p)

# ---------- 머리글 · 바닥글 ----------
sec.different_first_page_header_footer = True
hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run("Nooki 하드웨어 부품 정리 보고서"); hr.font.size = Pt(8); hr.font.color.rgb = RGBColor.from_string(SLATE)
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
def field(p, code):
    r = p.add_run(); r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(SLATE)
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); r._r.append(b)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = code; r._r.append(it)
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r._r.append(e)
field(fp, " PAGE ")
sr = fp.add_run(" / "); sr.font.size = Pt(8); sr.font.color.rgb = RGBColor.from_string(SLATE)
field(fp, " NUMPAGES ")

d.save(dst)
print("saved", dst)
